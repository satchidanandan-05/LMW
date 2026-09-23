import csv
import io
from datetime import datetime

import openpyxl

from services import audit_service
from services import export_service as ex
from services.auth_service import User
from services.report_service import generate_report, record_export, trace_chain
from services.traceability_service import get_traceability

USER = User(3, "supervisor1", "supervisor")
NOW = datetime(2026, 9, 23, 16, 18, 55)


def test_report_matches_search_and_is_audited(engine):
    rpt = generate_report(engine, "CY", "yc006", USER, now=NOW)
    assert rpt.trace == get_traceability(engine, "CY", "YC006")
    assert rpt.title == "Traceability Report - Yarn Cone YC006"
    assert (rpt.generated_by, rpt.generated_at) == ("supervisor1", NOW)
    row = audit_service.recent(engine)[0]
    assert rpt.reference_no == f"RPT-{row['audit_id']}"
    assert (row["action"], row["result"], row["id_value"]) == ("REPORT", "FOUND", "YC006")


def test_trace_chain_text(engine):
    assert trace_chain(get_traceability(engine, "CY", "YC006")) == (
        "YC006 → AC-02 / D025 → COB004 (SF-02/S128), COB005 (SF-02/S129), COB015 (SF-05/S210)"
    )
    assert trace_chain(get_traceability(engine, "CY", "YC006"), arrow="->").startswith("YC006 -> AC-02 / D025 -> ")
    assert trace_chain(get_traceability(engine, "CY", "YC007")) == "YC007 → AC-02 / D026 → no COBs"
    assert trace_chain(get_traceability(engine, "COB", "COB099")) == "COB099 (SF-02/S130) → no Yarn Cone link"
    assert trace_chain(get_traceability(engine, "CY", "YC999")) == "No traceability data"


def test_record_export_is_audited(engine):
    rpt = generate_report(engine, "COB", "COB004", USER, now=NOW)
    record_export(engine, USER.user_id, rpt, "PDF")
    row = audit_service.recent(engine)[0]
    assert (row["action"], row["id_type"], row["id_value"]) == ("EXPORT", "COB", "COB004")
    assert row["detail"] == f"format=PDF; ref={rpt.reference_no}"


def test_file_stem(engine):
    assert ex.file_stem(generate_report(engine, "CY", "YC006", USER, now=NOW)) == "traceability_YC006_20260923_161855"


def test_csv_one_row_per_cob(engine):
    data = ex.to_csv(generate_report(engine, "CY", "YC006", USER, now=NOW))
    rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert list(rows[0].keys()) == ex.CSV_COLUMNS
    assert [r["COB ID"] for r in rows] == ["COB004", "COB005", "COB015"]
    assert rows[0]["Searched ID"] == "YC006"
    assert rows[0]["Cone Scan Date & Time"] == "20-09-2026 11:12:09"


def test_csv_cone_without_cobs_has_single_row(engine):
    data = ex.to_csv(generate_report(engine, "CY", "YC007", USER, now=NOW))
    rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert len(rows) == 1
    assert rows[0]["Yarn Cone ID"] == "YC007" and rows[0]["COB ID"] == ""


def test_excel_sheets(engine):
    wb = openpyxl.load_workbook(io.BytesIO(ex.to_excel(generate_report(engine, "CY", "YC006", USER, now=NOW))))
    assert wb.sheetnames == ["Summary", "COBs", "Exceptions"]
    assert wb["COBs"].max_row == 4
    assert wb["COBs"]["B2"].value == "COB004"
    assert wb["Exceptions"]["A2"].value == "None"


def test_pdf_is_generated(engine):
    pdf = ex.to_pdf(generate_report(engine, "CY", "YC008", USER, now=NOW))
    assert pdf.startswith(b"%PDF")
    assert b"YC008" in pdf  # document title metadata
    assert len(pdf) > 1500
